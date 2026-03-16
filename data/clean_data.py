import json
import re
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "phase1.xlsx"
OUTPUT_FILE = BASE_DIR / "resources.geojson"

wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
print("Loaded sheets:", wb.sheetnames)


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(",", " ")
    if text.startswith("="):
        return None
    match = re.fullmatch(r"-?\d+(?:\.\d+)?", text)
    if match:
        return float(match.group())
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if match:
        return float(match.group())
    return None


def parse_two_numbers_from_text(value):
    if value is None:
        return None, None

    text = str(value)
    if text.startswith("="):
        return None, None

    nums = re.findall(r"-?\d+(?:\.\d+)?", text)
    if len(nums) >= 2:
        return float(nums[0]), float(nums[1])
    return None, None


def normalize_coords(x_value, y_value):
    x = parse_number(x_value)
    y = parse_number(y_value)

    if x is None and y is None:
        a, b = parse_two_numbers_from_text(x_value)
        if a is not None and b is not None:
            x, y = a, b

    if x is None or y is None:
        return None, None

    possible_pairs = [
        (x, y),
        (y, x),
    ]

    for lat, lng in possible_pairs:
        if 41 <= lat <= 43 and -84.5 <= lng <= -82:
            return lat, lng

    return None, None


def category_from_sheet(title):
    lowered = title.lower()
    if "community" in lowered:
        return "community"
    if "business" in lowered:
        return "business"
    if "transportation" in lowered:
        return "transportation"
    if "public" in lowered:
        return "public_resources"
    return "other"


def looks_like_header(row_values):
    joined = " ".join(clean_text(v).lower() for v in row_values if v is not None)
    return (
        "resource title" in joined
        or "source info" in joined
        or "x-coordinate" in joined
        or "y-coordinate" in joined
    )


def row_has_url(row_values):
    for value in row_values:
        text = clean_text(value).lower()
        if "http" in text or "www." in text:
            return True
    return False


def extract_row_fields(values):
    text_values = [clean_text(v) for v in values]

    name = ""
    source = ""
    description = ""
    lat = None
    lng = None

    for text in text_values:
        lowered = text.lower()
        if not source and ("http" in lowered or "www." in lowered):
            source = text

    for i, value in enumerate(values):
        a, b = parse_two_numbers_from_text(value)
        if a is not None:
            maybe_lat, maybe_lng = normalize_coords(a, b)
            if maybe_lat is not None:
                lat, lng = maybe_lat, maybe_lng
                break

    if lat is None:
        numeric_candidates = []
        for value in values:
            number = parse_number(value)
            if number is not None:
                numeric_candidates.append(number)

        for i in range(len(numeric_candidates) - 1):
            first = numeric_candidates[i]
            second = numeric_candidates[i + 1]
            maybe_lat, maybe_lng = normalize_coords(first, second)
            if maybe_lat is not None:
                lat, lng = maybe_lat, maybe_lng
                break

    for text in text_values:
        lowered = text.lower()
        if not text:
            continue
        if text == source:
            continue
        if lowered in {"resource title", "source info", "x-coordinate", "y-coordinate"}:
            continue
        if lowered.startswith("http") or lowered.startswith("www."):
            continue
        if parse_two_numbers_from_text(text)[0] is not None:
            continue
        if parse_number(text) is not None:
            continue
        name = text
        break

    leftover_text = []
    for text in text_values:
        lowered = text.lower()
        if not text:
            continue
        if text == name or text == source:
            continue
        if lowered.startswith("http") or lowered.startswith("www."):
            continue
        if parse_two_numbers_from_text(text)[0] is not None:
            continue
        if parse_number(text) is not None:
            continue
        leftover_text.append(text)

    if leftover_text:
        description = " | ".join(leftover_text)

    return name, source, description, lat, lng


features = []
debug_rows_checked = 0

for ws in wb.worksheets:
    category = category_from_sheet(ws.title)

    for row in ws.iter_rows(values_only=True):
        values = list(row)

        if not any(v is not None and str(v).strip() != "" for v in values):
            continue

        if looks_like_header(values):
            continue

        name, source, description, lat, lng = extract_row_fields(values)

        if debug_rows_checked < 12:
            print("SHEET:", ws.title)
            print("ROW:", values)
            print("PARSED:", {"name": name, "source": source, "lat": lat, "lng": lng})
            print("---")
            debug_rows_checked += 1

        if not name:
            continue

        if lat is None or lng is None:
            continue

        if not source and not row_has_url(values):
            continue

        features.append({
            "type": "Feature",
            "properties": {
                "name": name,
                "category": category,
                "sheet": ws.title,
                "source": source,
                "description": description,
            },
            "geometry": {
                "type": "Point",
                "coordinates": [lng, lat],
            },
        })

geojson = {"type": "FeatureCollection", "features": features}

with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
    json.dump(geojson, f, indent=2)

print(f"Wrote {len(features)} features to {OUTPUT_FILE}")